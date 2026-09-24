#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
insert_profile_photos.py
========================

Embeds each learner's photograph into the "Profile Picture" column of an
Excel workbook, matching rows to photos by the unique "App.No" value.

    Excel App.No : MSU-WI/2026-27/15432
    ZIP folder   : <any root>/MSU-WI-2026-27-15432/
    ZIP photo    : PhotoMSU-WI-2026-27-15432.jpg   (.jpg / .jpeg / .png)

* The ZIP is read directly - no manual extraction needed.
* The original Excel file is never modified; a new "_with_photos" file is made.
* A CSV report lists every record that could not be matched (and every
  duplicate), so nothing fails silently.

Quick start (Windows):
    1. pip install openpyxl pillow
    2. python insert_profile_photos.py
    3. A window opens: choose the Excel file, then choose the ZIP file.
       (You can also pass the two paths on the command line.)

How a student is matched
    Folder  : exact name -> ignoring - / _ spaces -> folder name contains the
              App.No -> folder name contains the serial number (e.g. 15432).
    Photo   : file starting with "Photo" -> file with "photo"/"profile"/"pic"
              in its name -> the only image in the folder that is not an
              Aadhar/ID/certificate scan. The App.No need NOT be in the
              file name, because the folder already identifies the student.

Requires: Python 3.8+, openpyxl, pillow
"""

from __future__ import annotations

import argparse
import csv
import io
import re
import sys
import zipfile
from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
    from openpyxl.utils.units import pixels_to_EMU
    from PIL import Image as PILImage
    from PIL import ImageOps
except ImportError as exc:  # pragma: no cover
    print("\nERROR: A required package is missing (%s)." % exc)
    print("Install the requirements with:\n")
    print("    pip install openpyxl pillow\n")
    sys.exit(1)


# =============================================================================
# SETTINGS  -  the only section a non-programmer might want to edit
# =============================================================================

# Default input file names (used when no file names are given on the command
# line). If these are not found, the script looks for the only .xlsx / .zip
# in the same folder as this script.
DEFAULT_EXCEL_NAME = "IEMS Learner Data Jan 2026.xlsx"
DEFAULT_ZIP_NAME = "IEMS Documents.zip"

# Output naming:  <excel name>_with_photos.xlsx  and  <excel name>_photo_report.csv
OUTPUT_SUFFIX = "_with_photos"
REPORT_SUFFIX = "_photo_report.csv"
FULL_REPORT_SUFFIX = "_full_report.xlsx"   # every student: all details + photo + status

# Column headers to look for (case, spaces and dots are ignored when matching).
APP_NO_HEADER = "App.No"
PICTURE_HEADER = "Profile Picture"

# Maximum photo size shown inside the cell, in screen pixels.
# The ratio 4:5 suits passport-style photographs.
MAX_PHOTO_WIDTH_PX = 112
MAX_PHOTO_HEIGHT_PX = 140

# Empty space (pixels) kept between the photo and the cell border.
CELL_PADDING_PX = 6

# Photos are shrunk before embedding to keep the Excel file small.
# They are stored at this multiple of the display size (2 = sharp on
# high-DPI screens). Photos are never enlarged beyond their original pixels.
EMBED_RESOLUTION_FACTOR = 2
JPEG_QUALITY = 85

# False = a small photo stays at its natural size instead of being blown up.
ALLOW_UPSCALE = False

# Vertically centre the text of each student row (rows become tall, and
# Excel would otherwise push the text to the bottom of the row).
# Set to False to leave every existing cell alignment completely untouched.
CENTER_ROW_CONTENT_VERTICALLY = True

# --- Photo file selection inside a student's folder -------------------------
# 1st choice: an image whose name STARTS with this text (e.g. PhotoMSU-...jpg)
PHOTO_FILENAME_PREFIX = "photo"
# 2nd choice: an image whose name CONTAINS one of these words
PHOTO_KEYWORDS = ("photo", "profile", "pic")
# 3rd choice: the only image in the folder, provided it does not look like a
# document scan (names starting with these words are ignored in this step)
DOCUMENT_WORDS = ("aadhar", "aadhaar", "adhar", "card", "marksheet", "mark",
                  "certificate", "signature", "proof", "degree", "migration",
                  "income", "caste", "cheque", "bank", "passbook")
DOCUMENT_EXACT_TOKENS = ("pan", "id", "sign", "sig")
ALLOWED_EXTENSIONS = {".jpg", ".jpeg", ".png"}

# --- Folder matching ----------------------------------------------------------
# True  = if no folder has exactly the App.No name, also accept folders whose
#         name only partly matches (contains the App.No, or contains the serial
#         number such as 15432). Every such match is listed in the report so
#         you can double-check it.
# False = exact folder names only.
ALLOW_PARTIAL_FOLDER_MATCH = True
MIN_SERIAL_DIGITS = 3      # the serial number must have at least this many digits

# How many top rows to scan when looking for the header row.
HEADER_SEARCH_ROWS = 30

# Characters allowed in an App.No (anything else is reported as invalid).
VALID_APP_NO_PATTERN = re.compile(r"^[A-Za-z0-9][A-Za-z0-9/_.\- ]*$")

# =============================================================================
# END OF SETTINGS
# =============================================================================


# Outcome codes ---------------------------------------------------------------
INSERTED = "INSERTED"
MISSING_FOLDER = "MISSING_FOLDER"
MISSING_PHOTO = "MISSING_PHOTO"
INVALID_APP_NO = "INVALID_APP_NO"
AMBIGUOUS_DUPLICATE = "AMBIGUOUS_DUPLICATE"
UNREADABLE_IMAGE = "UNREADABLE_IMAGE"


@dataclass
class FileEntry:
    zip_path: str          # exact name inside the ZIP
    dir_path: str          # folder containing the file
    filename: str
    crc: int
    stem_alnum: str = ""   # file name without extension, letters+digits only


@dataclass
class Record:
    row: int
    app_no: str
    key: str = ""
    alnum: str = ""
    outcome: str = ""
    dup_excel: bool = False
    dup_zip: bool = False
    partial_folder: bool = False
    fallback_photo: bool = False
    matched_folders: List[str] = field(default_factory=list)
    photo_file: str = ""            # file used inside the ZIP
    photo_bytes: bytes = b""        # resized JPEG that was embedded
    photo_size: Tuple[int, int] = (0, 0)
    issues: List[Tuple[str, str]] = field(default_factory=list)  # (code, detail)


@dataclass
class ZipIndex:
    by_key: Dict[str, List[str]]                 # 'MSU-WI-2026-27-15432' -> dir paths
    by_alnum: Dict[str, List[str]]               # 'MSUWI20262715432'     -> dir paths
    dir_info: Dict[str, Tuple[str, str]]         # dir path -> (name, alnum name)
    files_under: Dict[str, Dict[str, FileEntry]]  # dir path -> every file below it
    leaf_dirs: Set[str]                          # dirs that directly hold files
    all_images: List[FileEntry] = field(default_factory=list)  # every .jpg/.jpeg/.png


# -----------------------------------------------------------------------------
# Small helpers
# -----------------------------------------------------------------------------
def norm_header(value) -> str:
    """'App.No' / 'app no' / 'APP_NO' -> 'appno'"""
    return re.sub(r"[^a-z0-9]", "", str(value).lower()) if value is not None else ""


def clean_text(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\u00a0", " ").strip()


def make_key(text: str) -> str:
    """Turn an App.No or folder name into the common matching key.
    'MSU-WI/2026-27/15432'  ->  'MSU-WI-2026-27-15432'"""
    text = clean_text(text).replace("\\", "/").replace("/", "-")
    text = re.sub(r"\s+", "", text)
    return text.upper()


def alnum(text: str) -> str:
    """Letters+digits only, upper case: 'msu_wi 2026-27/15432' -> 'MSUWI20262715432'"""
    return re.sub(r"[^A-Z0-9]", "", clean_text(text).upper())


def last_serial(app_no: str) -> str:
    """Last group of digits in the App.No ('MSU-WI/2026-27/15432' -> '15432')."""
    groups = re.findall(r"\d+", app_no)
    if groups and len(groups[-1]) >= MIN_SERIAL_DIGITS:
        return groups[-1]
    return ""


def is_image_name(filename: str) -> bool:
    return Path(filename.lower()).suffix in ALLOWED_EXTENSIONS


def looks_like_document(filename: str) -> bool:
    """True for names like 'Aadhar CardMSU-...jpg' or 'PAN.png'."""
    stem = Path(filename.lower()).stem
    tokens = [t for t in re.split(r"[^a-z0-9]+", stem) if t]
    for t in tokens:
        if t in DOCUMENT_EXACT_TOKENS or any(t.startswith(w) for w in DOCUMENT_WORDS):
            return True
    return False


def col_width_for_pixels(px: int) -> float:
    """Excel stores column width so that pixels = width * 7 (default font)."""
    return px / 7.0


def row_height_for_pixels(px: int) -> float:
    """Excel row height is in points: 1 px = 0.75 pt."""
    return px * 0.75


# -----------------------------------------------------------------------------
# Input file discovery
# -----------------------------------------------------------------------------
def resolve_input(cli_value: Optional[str], default_name: str,
                  patterns: Tuple[str, ...], label: str,
                  script_dir: Path) -> Path:
    if cli_value:
        path = Path(cli_value.strip('"')).expanduser()
        if not path.is_absolute():
            path = Path.cwd() / path
        if not path.is_file():
            raise FileNotFoundError("%s not found: %s" % (label, path))
        return path

    default = script_dir / default_name
    if default.is_file():
        return default

    candidates = []
    for pattern in patterns:
        for p in script_dir.glob(pattern):
            if p.name.startswith("~$") or OUTPUT_SUFFIX in p.stem:
                continue
            candidates.append(p)
    if len(candidates) == 1:
        return candidates[0]
    if not candidates:
        raise FileNotFoundError(
            "%s not found. Put '%s' in the same folder as this script:\n    %s"
            % (label, default_name, script_dir))
    raise FileNotFoundError(
        "Several possible %s files found in %s:\n    %s\n"
        "Specify which one on the command line."
        % (label, script_dir, "\n    ".join(p.name for p in candidates)))


# -----------------------------------------------------------------------------
# Excel helpers
# -----------------------------------------------------------------------------
def locate_sheet_and_headers(wb):
    """Find the worksheet + header row containing the App.No column.
    Prefers a sheet that has both App.No and Profile Picture."""
    want_app = norm_header(APP_NO_HEADER)
    want_pic = norm_header(PICTURE_HEADER)
    fallback = None

    for ws in wb.worksheets:
        last = min(ws.max_row, HEADER_SEARCH_ROWS)
        for r in range(1, last + 1):
            cols: Dict[str, int] = {}
            for c in range(1, ws.max_column + 1):
                h = norm_header(ws.cell(r, c).value)
                if h and h not in cols:          # first occurrence wins
                    cols[h] = c
            if want_app in cols:
                found = (ws, r, cols[want_app], cols.get(want_pic))
                if want_pic in cols:
                    return found
                if fallback is None:
                    fallback = found
    if fallback:
        return fallback
    raise ValueError(
        "Could not find a column headed '%s' in any worksheet "
        "(searched the first %d rows of each sheet)."
        % (APP_NO_HEADER, HEADER_SEARCH_ROWS))


def row_is_blank(ws, row: int) -> bool:
    return all(clean_text(c.value) == "" for c in ws[row])


# -----------------------------------------------------------------------------
# ZIP helpers
# -----------------------------------------------------------------------------
def build_zip_index(zf: zipfile.ZipFile) -> ZipIndex:
    """Scan the ZIP once. Every folder level is indexed, so the name/depth of the
    root folder ('IEMS Documents/') does not matter, and a photo kept in a
    sub-folder of the student's folder is still found."""
    files_under: Dict[str, Dict[str, FileEntry]] = defaultdict(dict)
    dir_names: Dict[str, str] = {}
    leaf_dirs: Set[str] = set()
    all_images: List[FileEntry] = []

    for info in zf.infolist():
        if info.is_dir():
            continue
        path = info.filename.replace("\\", "/")           # Windows-made ZIPs
        parts = [p for p in path.split("/") if p]
        if not parts:
            continue
        if "__MACOSX" in parts or parts[-1].startswith("._") \
                or parts[-1].lower() == "thumbs.db":
            continue
        entry = FileEntry(info.filename, "/".join(parts[:-1]), parts[-1], info.CRC,
                          alnum(Path(parts[-1]).stem))
        if is_image_name(entry.filename):
            all_images.append(entry)
        if len(parts) < 2:              # file sitting directly in the ZIP root
            continue
        leaf_dirs.add(entry.dir_path)
        for i in range(1, len(parts)):
            d = "/".join(parts[:i])
            dir_names[d] = parts[i - 1]
            files_under[d][entry.zip_path] = entry

    by_key: Dict[str, List[str]] = defaultdict(list)
    by_alnum: Dict[str, List[str]] = defaultdict(list)
    dir_info: Dict[str, Tuple[str, str]] = {}
    for d, name in dir_names.items():
        by_key[make_key(name)].append(d)
        by_alnum[alnum(name)].append(d)
        dir_info[d] = (name, alnum(name))
    return ZipIndex(by_key, by_alnum, dir_info, files_under, leaf_dirs, all_images)


def top_level(paths: List[str]) -> List[str]:
    """Drop any folder that sits inside another matched folder."""
    chosen = set(paths)
    return sorted(p for p in chosen
                  if not any(p != q and p.startswith(q + "/") for q in chosen))


def find_student_folders(rec: Record, index: ZipIndex,
                         excel_alnums: Set[str]) -> Tuple[List[str], str]:
    """Return (folder paths, how it matched). Tries, in order:
       exact -> ignoring separators -> name contains App.No -> name contains serial.
    A folder that belongs exactly to a *different* Excel record is never taken."""
    paths = index.by_key.get(rec.key)
    if paths:
        return top_level(paths), "exact"
    paths = index.by_alnum.get(rec.alnum)
    if paths:
        return top_level(paths), "separator-insensitive"
    if not ALLOW_PARTIAL_FOLDER_MATCH:
        return [], ""

    others = excel_alnums - {rec.alnum}
    pattern = re.compile(r"(?<!\d)%s(?!\d)" % re.escape(rec.alnum))
    hits = [d for d, (name, na) in index.dir_info.items()
            if na not in others and pattern.search(na)]
    if hits:
        return top_level(hits), "folder name contains the App.No"

    serial = last_serial(rec.app_no)
    if serial:
        spattern = re.compile(r"(?<!\d)%s(?!\d)" % re.escape(serial))
        hits = [d for d, (name, na) in index.dir_info.items()
                if na not in others and spattern.search(name)]
        if hits:
            return top_level(hits), "folder name contains serial number %s" % serial
    return [], ""


def find_photo_files(rec: Record, index: ZipIndex,
                     excel_alnums: Set[str]) -> Tuple[List[FileEntry], str]:
    """For ZIPs where the photos are NOT in per-student folders (all photos in one
    folder, named like 'PhotoMSU-WI_2026-27_15526.jpg').
    Tries: file name contains the full App.No -> file name ends with the same
    serial number (last digits of the App.No, e.g. 15526)."""
    imgs = index.all_images
    pattern = re.compile(r"(?<!\d)%s(?!\d)" % re.escape(rec.alnum))
    hits = [f for f in imgs if pattern.search(f.stem_alnum)]
    if hits:
        return hits, "file name contains the App.No"
    if not ALLOW_PARTIAL_FOLDER_MATCH:
        return [], ""
    serial = last_serial(rec.app_no)
    if not serial:
        return [], ""
    others = [o for o in excel_alnums if o and o != rec.alnum]

    def belongs_to_other(f: FileEntry) -> bool:
        return any(re.search(r"(?<!\d)%s(?!\d)" % re.escape(o), f.stem_alnum) for o in others)

    def file_serial(f: FileEntry) -> str:
        groups = re.findall(r"\d+", Path(f.filename).stem)
        return groups[-1] if groups else ""

    hits = [f for f in imgs if file_serial(f) == serial and not belongs_to_other(f)]
    if hits:
        return hits, "file name has the same serial number %s" % serial
    return [], ""


def pick_photo_candidates(files: List[FileEntry]) -> Tuple[List[FileEntry], str]:
    """Choose the photograph among the files inside a student's folder.
    The App.No does NOT have to be in the file name - the folder identifies the student."""
    images = [f for f in files if is_image_name(f.filename)]
    if not images:
        return [], "no-image"
    tier1 = [f for f in images if f.filename.lower().startswith(PHOTO_FILENAME_PREFIX)]
    if tier1:
        return tier1, "photo-name"
    tier2 = [f for f in images if any(k in f.filename.lower() for k in PHOTO_KEYWORDS)]
    if tier2:
        return tier2, "photo-keyword"
    tier3 = [f for f in images if not looks_like_document(f.filename)]
    if tier3:
        return tier3, "only-image"
    return [], "only-documents"


# -----------------------------------------------------------------------------
# Image processing
# -----------------------------------------------------------------------------
def prepare_photo(data: bytes) -> Tuple[io.BytesIO, int, int]:
    """Return (jpeg_buffer, display_width_px, display_height_px).

    * respects EXIF rotation (phone photos)
    * flattens transparency onto white
    * scales proportionally to fit MAX_PHOTO_WIDTH_PX x MAX_PHOTO_HEIGHT_PX
    * stores a compact JPEG at EMBED_RESOLUTION_FACTOR x display size
    """
    with PILImage.open(io.BytesIO(data)) as im:
        im.load()
        im = ImageOps.exif_transpose(im)

        has_alpha = im.mode in ("RGBA", "LA") or \
            (im.mode == "P" and "transparency" in im.info)
        if has_alpha:
            rgba = im.convert("RGBA")
            background = PILImage.new("RGB", rgba.size, (255, 255, 255))
            background.paste(rgba, mask=rgba.split()[-1])
            im = background
        else:
            im = im.convert("RGB")

        w, h = im.size
        if w < 1 or h < 1:
            raise ValueError("image has zero size")

        scale = min(MAX_PHOTO_WIDTH_PX / w, MAX_PHOTO_HEIGHT_PX / h)
        if not ALLOW_UPSCALE:
            scale = min(scale, 1.0)
        disp_w = max(1, round(w * scale))
        disp_h = max(1, round(h * scale))

        # Pixel size actually stored in the workbook (never enlarge the source)
        store_scale = min(1.0,
                          (MAX_PHOTO_WIDTH_PX * EMBED_RESOLUTION_FACTOR) / w,
                          (MAX_PHOTO_HEIGHT_PX * EMBED_RESOLUTION_FACTOR) / h)
        if store_scale < 1.0:
            im = im.resize((max(1, round(w * store_scale)),
                            max(1, round(h * store_scale))),
                           PILImage.LANCZOS)

        buf = io.BytesIO()
        im.save(buf, format="JPEG", quality=JPEG_QUALITY, optimize=True)
        buf.seek(0)
        return buf, disp_w, disp_h


def add_centered_image(ws, buf: io.BytesIO, disp_w: int, disp_h: int,
                       row: int, col: int, cell_w_px: int, cell_h_px: int):
    """Place the image centred inside a single cell.

    A two-cell anchor ('move and size with cells') is used so the photo travels
    with its row when the sheet is filtered or rows are hidden."""
    img = XLImage(buf)
    img.width, img.height = disp_w, disp_h

    off_x = max(0, (cell_w_px - disp_w) // 2)
    off_y = max(0, (cell_h_px - disp_h) // 2)

    start = AnchorMarker(col=col - 1, colOff=pixels_to_EMU(off_x),
                         row=row - 1, rowOff=pixels_to_EMU(off_y))
    end = AnchorMarker(col=col - 1, colOff=pixels_to_EMU(off_x + disp_w),
                       row=row - 1, rowOff=pixels_to_EMU(off_y + disp_h))
    img.anchor = TwoCellAnchor(editAs="twoCell", _from=start, to=end)
    ws.add_image(img)


# -----------------------------------------------------------------------------
# Main workflow
# -----------------------------------------------------------------------------
def process(excel_path: Path, zip_path: Path) -> Tuple[int, str]:
    started = datetime.now()

    # ---- Load workbook ------------------------------------------------------
    print("Loading Excel : %s" % excel_path.name)
    keep_vba = excel_path.suffix.lower() == ".xlsm"
    wb = load_workbook(excel_path, keep_vba=keep_vba)
    ws, header_row, app_col, pic_col = locate_sheet_and_headers(wb)
    print("Worksheet     : '%s'  (header on row %d)" % (ws.title, header_row))

    if pic_col is None:
        pic_col = ws.max_column + 1
        header_cell = ws.cell(header_row, pic_col, PICTURE_HEADER)
        prev = ws.cell(header_row, pic_col - 1)
        if prev.has_style:
            header_cell._style = copy(prev._style)
        print("Note: no '%s' column existed, so one was added (column %d)."
              % (PICTURE_HEADER, pic_col))
    print("Columns       : App.No = %d, %s = %d" % (app_col, PICTURE_HEADER, pic_col))

    # ---- Read ZIP -----------------------------------------------------------
    print("Reading ZIP   : %s" % zip_path.name)
    zf = zipfile.ZipFile(zip_path, "r")
    try:
        index = build_zip_index(zf)
        print("ZIP contains  : %d folders holding files, %d files in total"
              % (len(index.leaf_dirs), len({e for d in index.leaf_dirs for e in index.files_under[d]})))

        # ---- Collect Excel records ------------------------------------------
        records: List[Record] = []
        for r in range(header_row + 1, ws.max_row + 1):
            if row_is_blank(ws, r):
                continue
            raw = clean_text(ws.cell(r, app_col).value)
            rec = Record(row=r, app_no=raw, key=make_key(raw), alnum=alnum(raw))
            records.append(rec)

        key_counts = Counter(rec.key for rec in records if rec.key)
        flat_dirs: Set[str] = set()      # folders whose photos were matched by file name
        used_files: Set[str] = set()     # photo files actually inserted
        excel_alnums = {rec.alnum for rec in records if rec.alnum}

        # ---- Cell geometry --------------------------------------------------
        cell_w = MAX_PHOTO_WIDTH_PX + 2 * CELL_PADDING_PX
        cell_h = MAX_PHOTO_HEIGHT_PX + 2 * CELL_PADDING_PX
        if row_height_for_pixels(cell_h) > 409:
            raise ValueError("MAX_PHOTO_HEIGHT_PX is too large: Excel rows "
                             "cannot be taller than 409 points.")
        ws.column_dimensions[
            ws.cell(1, pic_col).column_letter].width = col_width_for_pixels(cell_w)

        # ---- Process each record --------------------------------------------
        total = len(records)
        print("\nProcessing %d record(s)..." % total)
        for n, rec in enumerate(records, start=1):
            if n % 100 == 0:
                print("  ...%d / %d" % (n, total))

            # invalid App.No
            if not rec.app_no or not VALID_APP_NO_PATTERN.match(rec.app_no):
                rec.outcome = INVALID_APP_NO
                rec.issues.append((INVALID_APP_NO,
                                   "App.No is blank or contains invalid characters"))
                continue

            if key_counts[rec.key] > 1:
                rec.dup_excel = True
                rec.issues.append(("DUPLICATE_APP_NO_IN_EXCEL",
                                   "App.No appears %d times in the Excel sheet"
                                   % key_counts[rec.key]))

            folder_paths, how = find_student_folders(rec, index, excel_alnums)
            files: Dict[str, FileEntry] = {}
            if not folder_paths:
                # no student folder -> look for the photo by its FILE name instead
                flat, fhow = find_photo_files(rec, index, excel_alnums)
                if not flat:
                    rec.outcome = MISSING_FOLDER
                    serial = last_serial(rec.app_no)
                    rec.issues.append((MISSING_FOLDER,
                                       "No ZIP folder or photo file matches '%s'%s" % (
                                           rec.key,
                                           " (also tried serial number %s)" % serial
                                           if serial and ALLOW_PARTIAL_FOLDER_MATCH else "")))
                    continue
                files = {f.zip_path: f for f in flat}
                flat_dirs.update(f.dir_path for f in flat)
                rec.matched_folders = sorted({f.dir_path for f in flat if f.dir_path})
                if fhow != "file name contains the App.No":
                    rec.partial_folder = True
                    rec.issues.append(("PARTIAL_FILE_MATCH",
                                       "Matched because the %s -> %s (App.No in the file name differs; "
                                       "please verify)"
                                       % (fhow, " | ".join(f.zip_path for f in flat))))
            else:
                rec.matched_folders = folder_paths

                if how not in ("exact", "separator-insensitive"):
                    rec.partial_folder = True
                    rec.issues.append(("PARTIAL_FOLDER_MATCH",
                                       "Matched by %s -> %s (please verify)"
                                       % (how, " | ".join(folder_paths))))

                if len(folder_paths) > 1:
                    rec.dup_zip = True
                    rec.issues.append(("DUPLICATE_FOLDER_IN_ZIP",
                                       "Folder found %d times: %s"
                                       % (len(folder_paths), " | ".join(folder_paths))))

                for fp in folder_paths:
                    files.update(index.files_under.get(fp, {}))
            candidates, tier = pick_photo_candidates(list(files.values()))

            if not candidates:
                rec.outcome = MISSING_PHOTO
                detail = ("Folder '%s' has only document scans (Aadhar/ID/etc.), no photograph"
                          if tier == "only-documents"
                          else "Folder '%s' has no .jpg/.jpeg/.png image")
                rec.issues.append((MISSING_PHOTO, detail % (
                    folder_paths[0] if folder_paths else "matched files")))
                continue

            if len(candidates) > 1:
                rec.dup_zip = True
                names = " | ".join(c.zip_path for c in candidates)
                if len({c.crc for c in candidates}) > 1:
                    rec.outcome = AMBIGUOUS_DUPLICATE
                    if tier == "photo-name":
                        msg = "%d different 'Photo' files" % len(candidates)
                    else:
                        msg = "%d images and none is named 'Photo'" % len(candidates)
                    rec.issues.append(("MULTIPLE_DIFFERENT_PHOTOS",
                                       "%s - none inserted (check manually): %s"
                                       % (msg, names)))
                    continue
                rec.issues.append(("DUPLICATE_IDENTICAL_PHOTOS",
                                   "%d identical copies, first one used: %s"
                                   % (len(candidates), names)))

            chosen = candidates[0]
            if tier != "photo-name":
                rec.fallback_photo = True
                rec.issues.append(("PHOTO_FILE_NAME_FALLBACK",
                                   "No file starting with 'Photo'; used '%s' (%s) - please verify"
                                   % (chosen.zip_path,
                                      "name contains photo/profile/pic" if tier == "photo-keyword"
                                      else "only non-document image in the folder")))
            try:
                buf, w, h = prepare_photo(zf.read(chosen.zip_path))
                used_files.add(chosen.zip_path)
                rec.photo_file, rec.photo_bytes, rec.photo_size = \
                    chosen.zip_path, buf.getvalue(), (w, h)
                add_centered_image(ws, buf, w, h, rec.row, pic_col, cell_w, cell_h)
                rec.outcome = INSERTED
            except Exception as exc:                        # corrupt image etc.
                rec.outcome = UNREADABLE_IMAGE
                reason = ("file is corrupt or not a real image"
                          if isinstance(exc, PILImage.UnidentifiedImageError)
                          else "%s: %s" % (type(exc).__name__, exc))
                rec.issues.append((UNREADABLE_IMAGE,
                                   "Could not read '%s' (%s)" % (chosen.zip_path, reason)))

        # ---- Row height / alignment for all student rows --------------------
        for rec in records:
            ws.row_dimensions[rec.row].height = row_height_for_pixels(cell_h)
            if CENTER_ROW_CONTENT_VERTICALLY:
                for cell in ws[rec.row]:
                    al = copy(cell.alignment)
                    al.vertical = "center"
                    cell.alignment = al

        # ---- ZIP folders that no Excel row refers to -------------------------
        matched = {fp for rec in records for fp in rec.matched_folders}
        ancestors = {"/".join(m.split("/")[:i]) for m in matched
                     for i in range(1, len(m.split("/")))}
        orphan_folders = []
        for d in sorted(index.leaf_dirs):
            parts = d.split("/")
            inside_matched = any("/".join(parts[:i]) in matched
                                 for i in range(1, len(parts) + 1))
            if not inside_matched and d not in ancestors:
                orphan_folders.append(d)
        # photo files (matched by file name) that no Excel row uses
        for e in index.all_images:
            if e.dir_path in flat_dirs and e.zip_path not in used_files \
                    and not looks_like_document(e.filename):
                orphan_folders.append(e.zip_path)

        # ---- Save workbook (never overwrite the original) --------------------
        out_path = excel_path.with_name(excel_path.stem + OUTPUT_SUFFIX + excel_path.suffix)
        if out_path.resolve() == excel_path.resolve():
            raise RuntimeError("Refusing to overwrite the original file.")
        print("\nSaving workbook (this can take a moment for large files)...")
        try:
            wb.save(out_path)
        except PermissionError:
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            out_path = out_path.with_name("%s_%s%s" % (out_path.stem, stamp, out_path.suffix))
            print("Output file is open in Excel; saving as %s instead." % out_path.name)
            wb.save(out_path)
    finally:
        zf.close()

    # ---- Report CSV -----------------------------------------------------------
    report_path = excel_path.with_name(excel_path.stem + REPORT_SUFFIX)
    rows_written = write_report(report_path, records, orphan_folders)
    full_path = write_full_report(excel_path, ws, header_row, pic_col, records,
                                  orphan_folders, cell_w, cell_h)
    print("\nFull report (all details + photo): %s" % full_path)

    # ---- Summary ----------------------------------------------------------------
    return print_summary(records, orphan_folders, out_path, report_path, rows_written, started)


def write_report(path: Path, records: List[Record], orphans: List[str]) -> int:
    lines = []
    for rec in records:
        for code, detail in rec.issues:
            lines.append([rec.row, rec.app_no, rec.key, code,
                          "YES" if rec.outcome == INSERTED else "NO", detail])
    for key in orphans:
        if is_image_name(key):
            lines.append(["", "", key, "ZIP_PHOTO_NOT_IN_EXCEL", "NO",
                          "Photo file matches no Excel App.No"])
        else:
            lines.append(["", "", key, "ZIP_FOLDER_NOT_IN_EXCEL", "NO",
                          "Folder has files but matches no Excel App.No"])

    target = path
    try:
        fh = open(target, "w", newline="", encoding="utf-8-sig")
    except PermissionError:
        target = path.with_name("%s_%s%s" % (
            path.stem, datetime.now().strftime("%Y%m%d_%H%M%S"), path.suffix))
        fh = open(target, "w", newline="", encoding="utf-8-sig")
    with fh:
        writer = csv.writer(fh)
        writer.writerow(["Excel Row", "App.No", "Expected Folder Key",
                         "Issue", "Photo Inserted", "Details"])
        writer.writerows(lines)
    return len(lines)


def write_full_report(excel_path: Path, src_ws, header_row: int, pic_col: int,
                      records: List[Record], orphans: List[str],
                      cell_w: int, cell_h: int) -> Path:
    """Excel report with EVERY student: photo + all original columns + status."""
    wbr = Workbook()
    ws = wbr.active
    ws.title = "Full Report"
    src_cols = [c for c in range(1, src_ws.max_column + 1) if c != pic_col]
    extra = ["Photo Status", "Matched ZIP Folder", "Photo File Used", "Remarks"]
    headers = ["Photo"] + [clean_text(src_ws.cell(header_row, c).value) for c in src_cols] + extra

    head_fill = PatternFill("solid", fgColor="1F4E78")
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(1, i, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.column_dimensions["A"].width = col_width_for_pixels(cell_w)

    ok_fill = PatternFill("solid", fgColor="E2F0D9")
    bad_fill = PatternFill("solid", fgColor="F8D7DA")
    warn_fill = PatternFill("solid", fgColor="FFF2CC")

    for n, rec in enumerate(records, start=2):
        ws.row_dimensions[n].height = row_height_for_pixels(cell_h)
        for j, c in enumerate(src_cols, start=2):
            v = src_ws.cell(rec.row, c).value
            ws.cell(n, j, v).alignment = Alignment(vertical="center", wrap_text=True)
        base = 2 + len(src_cols)
        remarks = " | ".join("%s: %s" % (code, det) for code, det in rec.issues)
        status = "PHOTO INSERTED" if rec.outcome == INSERTED else (rec.outcome or "NOT PROCESSED")
        values = [status, " | ".join(rec.matched_folders), rec.photo_file, remarks]
        for k, v in enumerate(values):
            ws.cell(n, base + k, v).alignment = Alignment(vertical="center", wrap_text=True)
        colour = ok_fill if rec.outcome == INSERTED and not rec.issues else \
            (warn_fill if rec.outcome == INSERTED else bad_fill)
        ws.cell(n, base).fill = colour
        if rec.photo_bytes:
            add_centered_image(ws, io.BytesIO(rec.photo_bytes), rec.photo_size[0],
                               rec.photo_size[1], n, 1, cell_w, cell_h)

    for j in range(2, len(headers) + 1):
        letter = get_column_letter(j)
        longest = max([len(str(ws.cell(r, j).value or "")) for r in range(1, min(ws.max_row, 60) + 1)] + [8])
        ws.column_dimensions[letter].width = min(max(12, longest + 2), 45)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), max(ws.max_row, 2))

    # ---- Summary sheet --------------------------------------------------------
    sm = wbr.create_sheet("Summary", 0)
    count = Counter(r.outcome for r in records)
    rows = [("Total Excel records", len(records)),
            ("Photos inserted", count[INSERTED]),
            ("Missing folder in ZIP", count[MISSING_FOLDER]),
            ("Folder without photo", count[MISSING_PHOTO]),
            ("Duplicate / different photos skipped", count[AMBIGUOUS_DUPLICATE]),
            ("Invalid App.No", count[INVALID_APP_NO]),
            ("Unreadable images", count[UNREADABLE_IMAGE]),
            ("ZIP folders not in Excel", len(orphans))]
    sm.column_dimensions["A"].width = 40
    sm.column_dimensions["B"].width = 14
    for i, (label, val) in enumerate(rows, start=1):
        sm.cell(i, 1, label).font = Font(bold=True)
        sm.cell(i, 2, val)
    if orphans:
        sm.cell(len(rows) + 2, 1, "ZIP folders not found in Excel:").font = Font(bold=True)
        for i, o in enumerate(orphans, start=len(rows) + 3):
            sm.cell(i, 1, o)

    path = excel_path.with_name(excel_path.stem + FULL_REPORT_SUFFIX)
    try:
        wbr.save(path)
    except PermissionError:
        path = path.with_name("%s_%s%s" % (path.stem, datetime.now().strftime("%Y%m%d_%H%M%S"), path.suffix))
        wbr.save(path)
    return path


def print_summary(records, orphans, out_path, report_path, rows_written, started):
    count = Counter(rec.outcome for rec in records)
    dup_records = sum(1 for r in records if r.dup_excel or r.dup_zip)
    missing = count[MISSING_FOLDER] + count[MISSING_PHOTO]

    line = "=" * 64
    print("\n" + line)
    print("  SUMMARY")
    print(line)
    print("  Total Excel records ............ %d" % len(records))
    print("  Photos inserted ................ %d" % count[INSERTED])
    print("      via partial folder name .... %d  (see report, please verify)"
          % sum(1 for r in records if r.outcome == INSERTED and r.partial_folder))
    print("      photo not named 'Photo' .... %d  (see report, please verify)"
          % sum(1 for r in records if r.outcome == INSERTED and r.fallback_photo))
    print("  Photos missing ................. %d" % missing)
    print("      no folder in ZIP ........... %d" % count[MISSING_FOLDER])
    print("      folder without photo ....... %d" % count[MISSING_PHOTO])
    print("  Duplicate / multiple matches ... %d" % dup_records)
    print("      skipped (different photos) . %d" % count[AMBIGUOUS_DUPLICATE])
    print("  Invalid App.No values .......... %d" % count[INVALID_APP_NO])
    print("  Unreadable / corrupt images .... %d" % count[UNREADABLE_IMAGE])
    print("  ZIP folders not in Excel ....... %d" % len(orphans))
    print(line)
    accounted = (count[INSERTED] + missing + count[AMBIGUOUS_DUPLICATE]
                 + count[INVALID_APP_NO] + count[UNREADABLE_IMAGE])
    print("  Check: inserted + not inserted = %d (records = %d) %s"
          % (accounted, len(records), "OK" if accounted == len(records) else "MISMATCH!"))
    print(line)
    print("  Output workbook : %s" % out_path)
    if rows_written:
        print("  Report (%d line%s): %s" % (rows_written, "" if rows_written == 1 else "s", report_path))
    else:
        print("  Report          : %s  (no issues found)" % report_path)
    print("  Time taken      : %.1f s" % (datetime.now() - started).total_seconds())
    print(line)

    problems = [r for r in records if r.outcome not in (INSERTED, "")]
    if problems:
        print("\n  Records needing attention (first 20):")
        for rec in problems[:20]:
            print("   row %-5d %-28s %s" % (rec.row, rec.app_no or "(blank)", rec.outcome))
        if len(problems) > 20:
            print("   ... and %d more - see the report file." % (len(problems) - 20))

    popup = ("Photos inserted: %d of %d records\nMissing: %d | Duplicates: %d | Invalid App.No: %d\n\n"
             "Workbook:\n%s\n\nReport:\n%s"
             % (count[INSERTED], len(records), missing, dup_records,
                count[INVALID_APP_NO], out_path, report_path))
    return 0, popup


# -----------------------------------------------------------------------------
# File selection (upload) window
# -----------------------------------------------------------------------------
class UserCancelled(Exception):
    pass


def pick_file_dialog(title: str, filetypes, initial_dir: Path):
    """Open a standard Windows 'Open file' window.
    Returns (window_available, chosen_path_or_None)."""
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
    except Exception:
        return False, None              # no tkinter / no display
    try:
        root.withdraw()
        root.attributes("-topmost", True)
        chosen = filedialog.askopenfilename(title=title, filetypes=filetypes,
                                            initialdir=str(initial_dir))
    finally:
        root.destroy()
    return True, (Path(chosen) if chosen else None)


def show_popup(title: str, message: str) -> None:
    try:
        import tkinter as tk
        from tkinter import messagebox
        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        messagebox.showinfo(title, message)
        root.destroy()
    except Exception:
        pass


def ask_for_file(label: str, filetypes, default_name: str,
                 patterns: Tuple[str, ...], script_dir: Path) -> Path:
    print("Select the %s file in the window that opens..." % label)
    available, chosen = pick_file_dialog("Select the %s file" % label, filetypes, script_dir)
    if available:
        if chosen is None:
            raise UserCancelled("No %s file selected." % label)
        return chosen
    # No window possible: try the script's folder, then ask for a typed path.
    try:
        return resolve_input(None, default_name, patterns, label, script_dir)
    except FileNotFoundError:
        if sys.stdin and sys.stdin.isatty():
            typed = input("Paste the full path of the %s file: " % label).strip().strip('"')
            return resolve_input(typed, default_name, patterns, label, script_dir)
        raise


def check_inputs(excel_path: Path, zip_path: Path) -> None:
    if excel_path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValueError("'%s' is not an .xlsx file. If it is an old .xls file, open it "
                         "in Excel and use File > Save As > Excel Workbook (.xlsx)."
                         % excel_path.name)
    if not zipfile.is_zipfile(zip_path):
        raise ValueError("'%s' is not a valid ZIP file." % zip_path.name)


def parse_args(argv):
    p = argparse.ArgumentParser(
        description="Embed learner photographs from a ZIP into an Excel sheet by App.No.")
    p.add_argument("excel", nargs="?",
                   help="Input .xlsx (leave out to choose it in a window)")
    p.add_argument("zip", nargs="?",
                   help="Input .zip (leave out to choose it in a window)")
    return p.parse_args(argv)


def main(argv=None) -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")
        except Exception:
            pass

    args = parse_args(argv)
    script_dir = Path(__file__).resolve().parent
    exit_code = 1
    interactive = not (args.excel or args.zip)
    try:
        if interactive:
            excel_path = ask_for_file("Excel", (("Excel files", "*.xlsx *.xlsm"),
                                                ("All files", "*.*")),
                                      DEFAULT_EXCEL_NAME, ("*.xlsx", "*.xlsm"), script_dir)
            zip_path = ask_for_file("ZIP", (("ZIP files", "*.zip"), ("All files", "*.*")),
                                    DEFAULT_ZIP_NAME, ("*.zip",), script_dir)
        else:
            excel_path = resolve_input(args.excel, DEFAULT_EXCEL_NAME,
                                       ("*.xlsx", "*.xlsm"), "Excel file", script_dir)
            zip_path = resolve_input(args.zip, DEFAULT_ZIP_NAME,
                                     ("*.zip",), "ZIP file", script_dir)
        check_inputs(excel_path, zip_path)
        exit_code, popup = process(excel_path, zip_path)
        if interactive:
            show_popup("Photo insertion finished", popup)
        try:                                   # open the folder holding the results
            import os
            os.startfile(str(excel_path.parent))
        except Exception:
            pass
    except UserCancelled as exc:
        print("\nCancelled: %s" % exc)
    except FileNotFoundError as exc:
        print("\nERROR: %s" % exc)
    except zipfile.BadZipFile:
        print("\nERROR: The ZIP file is corrupt or is not a valid ZIP archive.")
    except PermissionError as exc:
        print("\nERROR: Permission denied (%s). Close the file in Excel and try again." % exc)
    except ValueError as exc:
        print("\nERROR: %s" % exc)
    except Exception as exc:                                  # last-resort net
        print("\nUNEXPECTED ERROR: %s: %s" % (type(exc).__name__, exc))
    finally:
        # Keep the window open when the script is double-clicked on Windows.
        if len(sys.argv) == 1 and sys.stdin and sys.stdin.isatty():
            try:
                input("\nPress Enter to close...")
            except EOFError:
                pass
    return exit_code


if __name__ == "__main__":
    sys.exit(main())
